"""Spreadsheet files in and out, and the upload store.

Uploads (CSV or XLSX) are parsed into sheets of row dicts (str, int, float or
None; Excel dates become ISO strings) and kept under state/uploads/<id>/.
The store also exposes every recent upload as long-format cells
(upload_id, sheet_name, row_no, col_name, cell_value, value_type), which the
API serves to Zetaris as the REST table `upload_cells`, so agents profile an
uploaded file with Zetaris SQL like any other source.

Outputs are written as XLSX (result sheet, summary sheets and an AgentOffice
sheet recording who approved which code), plus a CSV of the result sheet when
the upload was a CSV.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import uuid
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Optional

from app.config import ROOT
from app.engines.world_model.ids import utc_now
from app.tasks.runtime import is_blank, is_num, text

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20_000
MAX_COLUMNS = 60
MAX_SHEETS = 10
CELLS_UPLOADS = 8          # how many recent uploads the Zetaris cells endpoint serves
DEMO_DIR = ROOT / "data" / "sample" / "tasks"
_INT = re.compile(r"^-?(0|[1-9]\d{0,14})$")
_FLOAT = re.compile(r"^-?(0|[1-9]\d{0,14})\.\d{1,10}$")


class FileError(ValueError):
    """An upload we cannot use; the message says why in plain words."""


def _cell(v: Any) -> Any:
    if v is None or isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.time() == time(0) else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, str):
        return v if v.strip() else None
    try:
        return float(v)
    except (TypeError, ValueError):
        return str(v)


def _csv_cell(v: str) -> Any:
    s = v.strip()
    if s == "":
        return None
    if _INT.match(s):
        return int(s)
    if _FLOAT.match(s):
        return float(s)
    return v


def _headers(raw: list[Any]) -> list[str]:
    out: list[str] = []
    for i, h in enumerate(raw):
        name = " ".join(str(h).split()) if h is not None and str(h).strip() else f"Column {i + 1}"
        base, n = name, 2
        while name in out:
            name, n = f"{base} ({n})", n + 1
        out.append(name)
    return out


def _table(name: str, raw_rows: list[list[Any]]) -> Optional[dict]:
    rows = [r for r in raw_rows if any(v is not None and str(v).strip() != "" for v in r)]
    if not rows:
        return None
    width = max(len(r) for r in rows)
    if width > MAX_COLUMNS:
        raise FileError(f"sheet {name!r} has {width} columns; the limit is {MAX_COLUMNS}")
    header = _headers(list(rows[0]) + [None] * (width - len(rows[0])))
    body = rows[1:]
    if len(body) > MAX_ROWS:
        raise FileError(f"sheet {name!r} has {len(body)} rows; the limit is {MAX_ROWS}")
    records = [{h: (r[i] if i < len(r) else None) for i, h in enumerate(header)} for r in body]
    named = [i < len(rows[0]) and not is_blank(rows[0][i]) for i in range(width)]
    # drop columns with no header and no values (stray formatting in Excel)
    keep = [h for i, h in enumerate(header) if named[i] or any(not is_blank(rec[h]) for rec in records)]
    return {"name": name, "columns": keep, "rows": [{h: rec[h] for h in keep} for rec in records]}


def parse(filename: str, data: bytes) -> dict:
    """{"format": "csv" | "xlsx", "sheets": [{"name", "columns", "rows"}]}."""
    if len(data) > MAX_BYTES:
        raise FileError(f"the file is {len(data) / 1e6:.1f} MB; the limit is {MAX_BYTES // 1e6:.0f} MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = data.decode("cp1252", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(content[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        raw = [[_csv_cell(v) for v in row] for row in csv.reader(io.StringIO(content), dialect)]
        table = _table(Path(filename).stem[:31] or "Sheet1", raw)
        sheets = [table] if table else []
        fmt = "csv"
    elif suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # openpyxl raises many types for a bad file
            raise FileError(f"this is not a readable Excel workbook ({type(e).__name__})") from None
        sheets = []
        for ws in wb.worksheets[:MAX_SHEETS]:
            table = _table(ws.title, [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)])
            if table:
                sheets.append(table)
        wb.close()
        fmt = "xlsx"
    else:
        raise FileError("upload a .csv or .xlsx file")
    if not sheets:
        raise FileError("the file has no rows")
    return {"format": fmt, "sheets": sheets}


def sheet_columns(parsed: dict) -> dict[str, list[str]]:
    return {s["name"]: s["columns"] for s in parsed["sheets"]}


def sheet(parsed: dict, name: str) -> dict:
    return next(s for s in parsed["sheets"] if s["name"] == name)


def cells(upload_id: str, parsed: dict) -> list[dict]:
    """Long-format cells for Zetaris: every value as text plus its type."""
    out = []
    for s in parsed["sheets"]:
        for i, row in enumerate(s["rows"], start=1):
            for col in s["columns"]:
                v = row.get(col)
                kind = "blank" if is_blank(v) else "number" if is_num(v) else "text"
                out.append({"upload_id": upload_id, "sheet_name": s["name"], "row_no": i, "col_name": col,
                            "cell_value": "" if kind == "blank" else text(v), "value_type": kind})
    return out


class UploadStore:
    def __init__(self, state_dir: Path):
        self.dir = Path(state_dir) / "uploads"
        self.out_dir = Path(state_dir) / "outputs"

    def save(self, filename: str, data: bytes, *, upload_id: Optional[str] = None) -> dict:
        name = Path(filename).name or "upload.csv"
        parsed = parse(name, data)
        upload_id = upload_id or "up_" + uuid.uuid4().hex[:12]
        folder = self.dir / upload_id
        folder.mkdir(parents=True, exist_ok=True)
        (folder / ("original" + Path(name).suffix.lower())).write_bytes(data)
        meta = {"upload_id": upload_id, "filename": name, "format": parsed["format"], "uploaded_at": utc_now(),
                "sheets": [{"name": s["name"], "columns": s["columns"], "rows": len(s["rows"])}
                           for s in parsed["sheets"]]}
        (folder / "parsed.json").write_text(json.dumps(parsed, ensure_ascii=False), encoding="utf-8")
        (folder / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1), encoding="utf-8")
        return meta

    def meta(self, upload_id: str) -> Optional[dict]:
        path = self.dir / Path(upload_id).name / "meta.json"
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None

    def parsed(self, upload_id: str) -> Optional[dict]:
        path = self.dir / Path(upload_id).name / "parsed.json"
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None

    def touch(self, upload_id: str) -> None:
        """Mark an upload as recently used, so it stays among the uploads Zetaris serves."""
        path = self.dir / Path(upload_id).name / "meta.json"
        if path.exists():
            os.utime(path)

    def ensure_demos(self) -> list[dict]:
        """The bundled demo files, stored once under fixed ids (demo_<stem>)."""
        out = []
        for path in sorted(DEMO_DIR.glob("*.*")) if DEMO_DIR.exists() else []:
            if path.suffix.lower() not in (".csv", ".xlsx"):
                continue
            upload_id = f"demo_{path.stem}"
            out.append(self.meta(upload_id) or self.save(path.name, path.read_bytes(), upload_id=upload_id))
        return out

    def recent(self, limit: int = CELLS_UPLOADS) -> list[str]:
        if not self.dir.exists():
            return []
        folders = [p for p in self.dir.iterdir() if (p / "meta.json").exists()]
        folders.sort(key=lambda p: (p / "meta.json").stat().st_mtime, reverse=True)
        return [p.name for p in folders[:limit]]

    def all_cells(self, include: Optional[str] = None) -> list[dict]:
        """Cells of the recent uploads (and `include`, if given) for the Zetaris REST table."""
        ids = self.recent()
        if include and include not in ids:
            ids.append(include)
        rows: list[dict] = []
        for upload_id in ids:
            parsed = self.parsed(upload_id)
            if parsed:
                rows += cells(upload_id, parsed)
        return rows

    # ---------------------------------------------------------------- outputs

    def output_path(self, file_id: str) -> Optional[Path]:
        name = Path(file_id).name
        path = self.out_dir / name
        return path if path.exists() and path.parent == self.out_dir else None

    def write_outputs(self, *, task_id: str, filename: str, fmt: str, sheet_name: str, columns: list[str],
                      rows: list[dict], sheets: dict[str, list[dict]], provenance: dict) -> list[dict]:
        self.out_dir.mkdir(parents=True, exist_ok=True)
        stem = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(filename).stem)[:40] or "result"
        files = []
        xlsx_name = f"{task_id}_{stem}_filled.xlsx"
        write_xlsx(self.out_dir / xlsx_name, sheet_name, columns, rows, sheets, provenance)
        files.append({"file_id": xlsx_name, "format": "xlsx", "rows": len(rows)})
        if fmt == "csv":
            csv_name = f"{task_id}_{stem}_filled.csv"
            with (self.out_dir / csv_name).open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(columns)
                writer.writerows([[text(r.get(c)) for c in columns] for r in rows])
            files.append({"file_id": csv_name, "format": "csv", "rows": len(rows)})
        return files


def write_xlsx(path: Path, sheet_name: str, columns: list[str], rows: list[dict], sheets: dict[str, list[dict]],
               provenance: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2937")

    def table(ws, cols: list[str], records: list[dict]) -> None:
        ws.append(cols)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for r in records:
            ws.append([r.get(c) for c in cols])
        ws.freeze_panes = "A2"
        if records:
            ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(cols, start=1):
            width = max([len(str(c))] + [len(text(r.get(c))) for r in records[:500]])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    table(ws, columns, rows)
    for name, records in sheets.items():
        cols = list(records[0]) if records else []
        table(wb.create_sheet(name[:31]), cols, records)
    about = wb.create_sheet("AgentOffice")
    about.append(["Field", "Value"])
    for cell in about[1]:
        cell.font, cell.fill = head_font, head_fill
    for key, value in provenance.items():
        about.append([key, value if isinstance(value, (str, int, float)) or value is None
                      else json.dumps(value, ensure_ascii=False)])
    about.column_dimensions["A"].width = 26
    about.column_dimensions["B"].width = 100
    for row in about.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
