"""Generate the demo spreadsheets for task mode (synthetic, seeded, reproducible).

    python scripts/generate_task_samples.py

Writes data/sample/tasks/:
  expense_claims.xlsx      claims with messy categories, mixed date formats and blank Team/Role columns;
                           the emails belong to the synthetic staff in data/sample/workforce/people.csv
                           (which Zetaris serves as workforce.people), plus two that are not on staff
  vendor_invoices.csv      an invoice register with duplicate invoice numbers and missing statuses
  leads_and_contacts.xlsx  a Leads sheet and a CRM Contacts sheet to check them against
  uae_site_claims.xlsx     bilingual claims whose staff IDs mix Western, Arabic-Indic and Persian digits
                           (P003 / P٠٠٣ / P۰۰۳), so a plain lookup on workforce.people misses them

Every person, email and company is invented; emails use the reserved example.com/example.org domains.
"""

from __future__ import annotations

import csv
import random
import sys
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "sample" / "tasks"
SEED = 20261025


def _people() -> list[dict]:
    with (ROOT / "data" / "sample" / "workforce" / "people.csv").open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _messy_email(email: str, rng: random.Random) -> str:
    roll = rng.random()
    if roll < 0.2:
        return email.upper()
    if roll < 0.35:
        return f" {email} "
    if roll < 0.45:
        return email.title()
    return email


def expense_claims(rng: random.Random) -> list[list]:
    people = _people()
    categories = ["travel", "Travel ", " TRAVEL", "meals", "Meals", "MEALS  ", "office supplies",
                  "Office  Supplies", "software", "Software", "training"]
    header = ["Claim ID", "Employee Email", "Date", "Category", "Description", "Amount", "Team", "Role"]
    rows = [header]
    start = date(2026, 9, 1)
    for i in range(1, 41):
        person = rng.choice(people)
        email = person["email"]
        if i in (17, 33):
            email = rng.choice(["temp.contractor@example.org", "jo.visitor@example.org"])
        day = start + timedelta(days=rng.randint(0, 40))
        style = rng.random()
        when = (day.strftime("%Y-%m-%d") if style < 0.4 else day.strftime("%d/%m/%Y") if style < 0.75
                else f"{day.day} {day.strftime('%b %Y')}")
        category = rng.choice(categories)
        amount = round(rng.uniform(12, 900), 2) if i != 25 else None
        description = {"travel": "Taxi to client site", "meals": "Team lunch", "office supplies": "Printer paper",
                       "software": "Monthly licence", "training": "Online course"}[" ".join(category.split()).lower()]
        rows.append([f"EC-{1000 + i}", _messy_email(email, rng), when, category, description, amount, None, None])
    return rows


def vendor_invoices(rng: random.Random) -> list[list]:
    vendors = ["Acme  Office Supply", "Northwind Logistics", "Blue Harbor IT", "Globex Facilities",
               "Initech Software", "Umbrella Catering"]
    header = ["Invoice No", "Vendor", "Invoice Date", "Due Date", "Amount", "Status"]
    rows = [header]
    seen = []
    for i in range(1, 31):
        issued = date(2026, 9, 1) + timedelta(days=rng.randint(0, 35))
        due = issued + timedelta(days=rng.choice([14, 30]))
        status = rng.choice(["Paid", "Paid", "Pending", "", "Approved", ""])
        row = [f"INV-{2400 + i}", rng.choice(vendors), issued.isoformat(), due.isoformat(),
               round(rng.uniform(80, 7500), 2), status]
        rows.append(row)
        seen.append(row)
    for original in rng.sample(seen, 4):   # re-entered duplicates, typed slightly differently
        dup = list(original)
        dup[0] = rng.choice([dup[0].lower(), f" {dup[0]}", f"{dup[0]} "])
        dup[5] = ""
        rows.insert(rng.randint(2, len(rows)), dup)
    return rows


def leads_and_contacts(rng: random.Random) -> tuple[list[list], list[list]]:
    first = ["Avery", "Blake", "Casey", "Devon", "Emerson", "Finley", "Harper", "Jordan", "Kendall", "Logan",
             "Morgan", "Parker", "Quinn", "Reese", "Rowan", "Sage", "Taylor", "Wren"]
    last = ["Ashford", "Bramley", "Carrow", "Dunmore", "Ellery", "Fairley", "Garrick", "Holloway", "Ingram",
            "Kestrel", "Langley", "Marlow"]
    companies = ["Contoso", "Fabrikam", "Tailspin Toys", "Wide World Importers", "Proseware", "Litware"]
    owners = ["Sales - North", "Sales - South", "Partnerships"]
    leads = [["Lead ID", "Full Name", "Email", "Company", "Phone"]]
    contacts = [["Contact ID", "Email", "Company", "Owner"]]
    people = []
    for i in range(1, 26):
        f, l = rng.choice(first), rng.choice(last)
        email = f"{f}.{l}{i}@example.com".lower()
        people.append((f, l, email))
        name = f"{f} {l}" if rng.random() > 0.15 else f"  {f}   {l} "
        leads.append([f"L-{300 + i}", name, _messy_email(email, rng), rng.choice(companies),
                      f"+61 4{rng.randint(10, 99)} {rng.randint(100, 999)} {rng.randint(100, 999)}"])
    for j, (f, l, email) in enumerate(rng.sample(people, 12), start=1):
        contacts.append([f"C-{700 + j}", email, rng.choice(companies), rng.choice(owners)])
    dup_email = contacts[3][1]
    contacts.append([f"C-{790}", dup_email.upper(), "Contoso", rng.choice(owners)])   # a duplicate CRM record
    return leads, contacts


ARABIC_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
PERSIAN_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def uae_site_claims(rng: random.Random) -> list[list]:
    """Claims from a UAE site: the same staff IDs typed in Western, Arabic-Indic and Persian digits."""
    people = _people()
    items = [("Taxi to client site", "تاكسي إلى موقع العميل"), ("Team lunch", "غداء الفريق"),
             ("Printer paper", "ورق طابعة"), ("Site visit parking", "موقف زيارة الموقع"),
             ("Online course", "دورة تدريبية")]
    header = ["Claim ID", "Staff ID", "Date", "Description", "الوصف", "Amount (AED)", "Team", "Role"]
    rows = [header]
    for i in range(1, 19):
        staff_id = rng.choice(people)["person_id"]
        style = i % 3 if i not in (7, 15) else 3
        if style == 1:
            staff_id = staff_id.translate(ARABIC_DIGITS)
        elif style == 2 and i % 2:
            staff_id = staff_id.translate(PERSIAN_DIGITS)
        elif style == 3:
            staff_id = f" {staff_id.lower()}"
        if i == 12:
            staff_id = "P٠٩٩"                    # not on staff
        english, arabic = rng.choice(items)
        day = date(2026, 10, 1) + timedelta(days=rng.randint(0, 20))
        rows.append([f"UAE-{500 + i}", staff_id, day.isoformat(), english, arabic,
                     round(rng.uniform(40, 2400), 2), None, None])
    return rows


def _xlsx(path: Path, sheets: dict[str, list[list]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
    wb.save(path)


def main() -> int:
    rng = random.Random(SEED)
    OUT.mkdir(parents=True, exist_ok=True)
    _xlsx(OUT / "expense_claims.xlsx", {"Claims": expense_claims(rng)})
    with (OUT / "vendor_invoices.csv").open("w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(vendor_invoices(rng))
    leads, contacts = leads_and_contacts(rng)
    _xlsx(OUT / "leads_and_contacts.xlsx", {"Leads": leads, "CRM Contacts": contacts})
    _xlsx(OUT / "uae_site_claims.xlsx", {"Claims": uae_site_claims(random.Random(SEED + 1))})
    for p in sorted(OUT.iterdir()):
        print(f"wrote {p.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
